import csv
import io
from datetime import datetime, date, time, timedelta
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session
from sqlalchemy import func, desc
from app.models.room import Room
from app.models.booking import Booking
from app.models.user import User, Department
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def resolve_date_range(range_type: str, start_date: Optional[date], end_date: Optional[date]):
    now = datetime.utcnow()
    if range_type == "today":
        start = datetime.combine(now.date(), time(0, 0, 0))
        end = datetime.combine(now.date(), time(23, 59, 59))
    elif range_type == "week":
        s_date = now.date() - timedelta(days=now.weekday())
        start = datetime.combine(s_date, time(0, 0, 0))
        end = datetime.combine(now.date() + timedelta(days=6 - now.weekday()), time(23, 59, 59))
    elif range_type == "month":
        start = datetime(now.year, now.month, 1, 0, 0, 0)
        # End of month
        if now.month == 12:
            end = datetime(now.year + 1, 1, 1, 23, 59, 59) - timedelta(days=1)
        else:
            end = datetime(now.year, now.month + 1, 1, 23, 59, 59) - timedelta(days=1)
    elif range_type == "year":
        start = datetime(now.year, 1, 1, 0, 0, 0)
        end = datetime(now.year, 12, 31, 23, 59, 59)
    elif range_type == "custom" and start_date and end_date:
        start = datetime.combine(start_date, time(0, 0, 0))
        end = datetime.combine(end_date, time(23, 59, 59))
    else:
        start = datetime(now.year, now.month, 1, 0, 0, 0)
        end = datetime.combine(now.date(), time(23, 59, 59))

    return start, end

def generate_report_summary(db: Session, start_time: datetime, end_time: datetime) -> Dict[str, Any]:
    bookings = db.query(Booking).filter(
        Booking.start_time >= start_time,
        Booking.start_time <= end_time
    ).all()

    total_bookings = len(bookings)
    cancelled_bookings = sum(1 for b in bookings if b.status == "cancelled")
    approved_bookings = sum(1 for b in bookings if b.status in ["approved", "completed"])
    pending_bookings = sum(1 for b in bookings if b.status == "pending")

    total_seconds = sum(
        (b.end_time - b.start_time).total_seconds()
        for b in bookings if b.status in ["approved", "completed"]
    )
    total_hours = round(total_seconds / 3600.0, 1)

    cancellation_rate = round((cancelled_bookings / total_bookings * 100), 1) if total_bookings > 0 else 0.0

    # Room stats
    all_rooms = db.query(Room).filter(Room.is_active == True).all()
    room_stats = []
    for r in all_rooms:
        r_bookings = [b for b in bookings if b.room_id == r.id and b.status in ["approved", "completed"]]
        r_hours = round(sum((b.end_time - b.start_time).total_seconds() for b in r_bookings) / 3600.0, 1)
        room_stats.append({
            "room_id": r.id,
            "room_name": r.name,
            "room_code": r.code,
            "capacity": r.capacity,
            "total_bookings": len(r_bookings),
            "total_hours": r_hours,
            "utilization_rate": round(r_hours / (30 * 8) * 100, 1)  # approx monthly base
        })

    room_stats.sort(key=lambda x: x["total_bookings"], reverse=True)
    most_used = room_stats[0] if room_stats and room_stats[0]["total_bookings"] > 0 else None
    least_used = room_stats[-1] if room_stats and room_stats[-1]["total_bookings"] >= 0 else None

    # Group by user
    user_counts: Dict[int, Dict[str, Any]] = {}
    for b in bookings:
        if b.user_id not in user_counts:
            dept_name = b.user.department.name if b.user and b.user.department else "ไม่ระบุ"
            user_counts[b.user_id] = {
                "user_id": b.user_id,
                "user_name": b.user.name if b.user else "ผู้ใช้งาน",
                "department_name": dept_name,
                "booking_count": 0,
                "total_hours": 0.0
            }
        user_counts[b.user_id]["booking_count"] += 1
        if b.status in ["approved", "completed"]:
            user_counts[b.user_id]["total_hours"] += round((b.end_time - b.start_time).total_seconds() / 3600.0, 1)

    user_stats = sorted(list(user_counts.values()), key=lambda x: x["booking_count"], reverse=True)

    status_distribution = {
        "approved": approved_bookings,
        "pending": pending_bookings,
        "cancelled": cancelled_bookings,
        "rejected": sum(1 for b in bookings if b.status == "rejected"),
        "completed": sum(1 for b in bookings if b.status == "completed")
    }

    return {
        "start_date": start_time.strftime("%Y-%m-%d"),
        "end_date": end_time.strftime("%Y-%m-%d"),
        "total_bookings": total_bookings,
        "total_hours": total_hours,
        "cancelled_bookings": cancelled_bookings,
        "cancellation_rate": cancellation_rate,
        "approved_bookings": approved_bookings,
        "pending_bookings": pending_bookings,
        "most_used_room": most_used,
        "least_used_room": least_used,
        "room_stats": room_stats,
        "user_stats": user_stats,
        "status_distribution": status_distribution
    }

def export_to_excel(summary_data: Dict[str, Any], bookings_list: List[Booking]) -> io.BytesIO:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"

    # Header style
    header_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    subhead_font = Font(name="Arial", size=11, bold=True, color="0F172A")
    table_head_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    table_head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    ws_summary["A1"] = "รายงานสรุปสถิติการใช้ห้องประชุม (Meeting Room Booking Summary)"
    ws_summary["A1"].font = header_font
    ws_summary["A1"].fill = header_fill
    ws_summary.merge_cells("A1:E1")

    ws_summary["A3"] = f"ช่วงเวลา: {summary_data['start_date']} ถึง {summary_data['end_date']}"
    ws_summary["A3"].font = subhead_font

    metrics = [
        ("การจองทั้งหมด", summary_data["total_bookings"]),
        ("ชั่วโมงการใช้งานรวม (ชั่วโมง)", summary_data["total_hours"]),
        ("การจองที่อนุมัติ", summary_data["approved_bookings"]),
        ("การจองที่รออนุมัติ", summary_data["pending_bookings"]),
        ("การจองที่ยกเลิก", summary_data["cancelled_bookings"]),
        ("อัตราการยกเลิก (%)", f"{summary_data['cancellation_rate']}%"),
    ]

    r = 5
    for label, val in metrics:
        ws_summary.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=r, column=2, value=val)
        r += 1

    # Room statistics table
    r += 2
    ws_summary.cell(row=r, column=1, value="สถิติแยกตามห้องประชุม").font = subhead_font
    r += 1
    room_cols = ["รหัสห้อง", "ชื่อห้องประชุม", "ความจุ (ที่นั่ง)", "จำนวนการจอง", "ชั่วโมงรวม"]
    for col_idx, col_name in enumerate(room_cols, 1):
        c = ws_summary.cell(row=r, column=col_idx, value=col_name)
        c.fill = table_head_fill
        c.font = table_head_font

    for item in summary_data["room_stats"]:
        r += 1
        ws_summary.cell(row=r, column=1, value=item["room_code"])
        ws_summary.cell(row=r, column=2, value=item["room_name"])
        ws_summary.cell(row=r, column=3, value=item["capacity"])
        ws_summary.cell(row=r, column=4, value=item["total_bookings"])
        ws_summary.cell(row=r, column=5, value=item["total_hours"])

    # Sheet 2: Bookings List
    ws_bookings = wb.create_sheet(title="Bookings Details")
    booking_cols = ["รหัสการจอง", "ชื่อห้อง", "หัวข้อการประชุม", "ผู้จอง", "เริ่ม", "สิ้นสุด", "จำนวนคน", "สถานะ"]
    for col_idx, col_name in enumerate(booking_cols, 1):
        c = ws_bookings.cell(row=1, column=col_idx, value=col_name)
        c.fill = table_head_fill
        c.font = table_head_font

    for row_idx, b in enumerate(bookings_list, 2):
        ws_bookings.cell(row=row_idx, column=1, value=b.booking_code)
        ws_bookings.cell(row=row_idx, column=2, value=b.room.name if b.room else "")
        ws_bookings.cell(row=row_idx, column=3, value=b.title)
        ws_bookings.cell(row=row_idx, column=4, value=b.user.name if b.user else b.organizer_name)
        ws_bookings.cell(row=row_idx, column=5, value=b.start_time.strftime("%Y-%m-%d %H:%M"))
        ws_bookings.cell(row=row_idx, column=6, value=b.end_time.strftime("%Y-%m-%d %H:%M"))
        ws_bookings.cell(row=row_idx, column=7, value=b.attendees_count)
        ws_bookings.cell(row=row_idx, column=8, value=b.status)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def export_to_csv(bookings_list: List[Booking]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Booking Code", "Room", "Title", "Organizer", "Start Time", "End Time", "Attendees", "Status"])
    for b in bookings_list:
        writer.writerow([
            b.booking_code,
            b.room.name if b.room else "",
            b.title,
            b.organizer_name,
            b.start_time.strftime("%Y-%m-%d %H:%M"),
            b.end_time.strftime("%Y-%m-%d %H:%M"),
            b.attendees_count,
            b.status
        ])
    return output.getvalue()

def export_to_pdf(summary_data: Dict[str, Any], bookings_list: List[Booking]) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=14
    )

    h2_style = ParagraphStyle(
        'H2Style',
        parent=styles['Heading2'],
        fontSize=13,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=10,
        spaceAfter=6
    )

    elements.append(Paragraph("Meeting Room Booking - Executive Report", title_style))
    elements.append(Paragraph(f"Period: {summary_data['start_date']} to {summary_data['end_date']}", styles['Normal']))
    elements.append(Spacer(1, 14))

    # Summary Table
    summary_table_data = [
        ["Total Bookings", str(summary_data["total_bookings"]), "Total Hours Used", f"{summary_data['total_hours']} hrs"],
        ["Approved Bookings", str(summary_data["approved_bookings"]), "Cancellation Rate", f"{summary_data['cancellation_rate']}%"],
        ["Pending Bookings", str(summary_data["pending_bookings"]), "Cancelled Bookings", str(summary_data["cancelled_bookings"])]
    ]
    t = Table(summary_table_data, colWidths=[130, 120, 130, 120])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#0F172A')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 14))

    # Room Statistics Table
    elements.append(Paragraph("Room Usage Breakdown", h2_style))
    room_table_data = [["Room Code", "Room Name", "Capacity", "Bookings", "Total Hours"]]
    for item in summary_data["room_stats"]:
        room_table_data.append([
            item["room_code"],
            item["room_name"],
            str(item["capacity"]),
            str(item["total_bookings"]),
            f"{item['total_hours']} h"
        ])

    rt = Table(room_table_data, colWidths=[90, 180, 70, 70, 90])
    rt.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    elements.append(rt)

    doc.build(elements)
    buffer.seek(0)
    return buffer
